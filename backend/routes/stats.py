"""Tableau de bord, visiteurs et rapports annuels (Excel).

- Les visites sont comptées une fois par jour et par navigateur (identifiant anonyme).
  Le navigateur de l'admin n'envoie rien (voir frontend/src/lib/visits.js).
- Les devis et commandes faits avec les courriels de test de la propriétaire ne sont pas comptés.
"""
from datetime import datetime, timedelta
from io import BytesIO
import logging
import re

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from auth import require_admin
from database import get_database

logger = logging.getLogger(__name__)
router = APIRouter(tags=["stats"])

# Courriels utilisés pour tester le site : exclus des chiffres et des rapports
DEFAULT_TEST_EMAILS = ["abrisia0plan@gmail.com", "noemie66667@gmail.com", "admin@abrisia-plan.ca"]

MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
        "août", "septembre", "octobre", "novembre", "décembre"]


class Visit(BaseModel):
    visitor_id: str
    path: str = "/"


async def _test_emails(db) -> list:
    doc = await db.site_settings.find_one({"key": "test_emails"})
    emails = (doc or {}).get("emails") or DEFAULT_TEST_EMAILS
    return [e.strip().lower() for e in emails if e and e.strip()]


def _not_test(field: str, emails: list) -> dict:
    """Filtre Mongo : exclut les courriels de test (sans tenir compte des majuscules)."""
    if not emails:
        return {}
    return {field: {"$not": re.compile("^(" + "|".join(re.escape(e) for e in emails) + ")$", re.I)}}


def _as_date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.replace("Z", ""))
        except ValueError:
            return None
    return None


@router.post("/visit")
async def record_visit(visit: Visit):
    """Compte un visiteur par jour (public, anonyme)."""
    vid = (visit.visitor_id or "")[:64]
    if not vid:
        return {"success": False}
    try:
        db = get_database()
        day = datetime.utcnow().strftime("%Y-%m-%d")
        await db.visits.update_one(
            {"day": day, "visitor_id": vid},
            {"$setOnInsert": {"day": day, "visitor_id": vid, "path": visit.path[:200], "created_at": datetime.utcnow()}},
            upsert=True,
        )
    except Exception as e:
        logger.warning(f"Visite non enregistrée: {e}")
    return {"success": True}


@router.get("/admin/dashboard")
async def dashboard(current_user: dict = Depends(require_admin)):
    db = get_database()
    emails = await _test_emails(db)
    now = datetime.utcnow()
    month_start = datetime(now.year, now.month, 1)
    year_start = datetime(now.year, 1, 1)

    devis_filter = _not_test("email", emails)
    orders_filter = _not_test("customer_email", emails)

    devis_mois = await db.devis.count_documents({**devis_filter, "created_at": {"$gte": month_start}})
    devis_attente = await db.devis.count_documents({**devis_filter, "status": "En attente"})
    commandes_mois = await db.kit_orders.count_documents({**orders_filter, "created_at": {"$gte": month_start}})
    commandes_annee = await db.kit_orders.count_documents({**orders_filter, "created_at": {"$gte": year_start}})
    revenus_annee = 0.0
    async for o in db.kit_orders.find({**orders_filter, "created_at": {"$gte": year_start},
                                        "status": {"$in": ["paid", "completed"]}}, {"total_amount": 1}):
        revenus_annee += float(o.get("total_amount") or 0)

    today = now.strftime("%Y-%m-%d")
    month_prefix = now.strftime("%Y-%m")
    visiteurs_jour = len(await db.visits.distinct("visitor_id", {"day": today}))
    visiteurs_mois = len(await db.visits.distinct("visitor_id", {"day": {"$regex": f"^{month_prefix}"}}))
    since = (now - timedelta(days=29)).strftime("%Y-%m-%d")
    visites_30j = await db.visits.count_documents({"day": {"$gte": since}})

    years = set()
    for coll in (db.devis, db.kit_orders):
        async for d in coll.find({}, {"created_at": 1}):
            dt = _as_date(d.get("created_at"))
            if dt:
                years.add(dt.year)
    years.add(now.year)

    return {
        "success": True,
        "mois": f"{MOIS[now.month - 1]} {now.year}",
        "annee": now.year,
        "devis_mois": devis_mois,
        "devis_attente": devis_attente,
        "commandes_mois": commandes_mois,
        "commandes_annee": commandes_annee,
        "revenus_annee": round(revenus_annee, 2),
        "visiteurs_jour": visiteurs_jour,
        "visiteurs_mois": visiteurs_mois,
        "visites_30j": visites_30j,
        "annees": sorted(years, reverse=True),
        "courriels_test": emails,
    }


@router.get("/admin/reports/{year}")
async def yearly_report(year: int, current_user: dict = Depends(require_admin)):
    """Rapport Excel d'une année : résumé par mois, devis, commandes."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="Module Excel non installé sur le serveur")

    db = get_database()
    emails = await _test_emails(db)
    start, end = datetime(year, 1, 1), datetime(year + 1, 1, 1)

    devis = [d async for d in db.devis.find({**_not_test("email", emails), "created_at": {"$gte": start, "$lt": end}})
             .sort("created_at", 1)]
    orders = [o async for o in db.kit_orders.find({**_not_test("customer_email", emails), "created_at": {"$gte": start, "$lt": end}})
              .sort("created_at", 1)]
    visits = [v async for v in db.visits.find({"day": {"$regex": f"^{year}-"}}, {"day": 1, "visitor_id": 1})]

    wb = Workbook()
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0F5E5A")

    def header(ws, cols, widths):
        ws.append(cols)
        for i, c in enumerate(ws[1]):
            c.font, c.fill = bold, fill
            ws.column_dimensions[c.column_letter].width = widths[i] if i < len(widths) else 18
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Résumé"
    header(ws, ["Mois", "Demandes de devis", "Commandes Collection", "Revenus payés ($)", "Visiteurs"], [14, 18, 20, 18, 12])
    for m in range(1, 13):
        dv = sum(1 for d in devis if _as_date(d.get("created_at")) and _as_date(d["created_at"]).month == m)
        od = [o for o in orders if _as_date(o.get("created_at")) and _as_date(o["created_at"]).month == m]
        rev = sum(float(o.get("total_amount") or 0) for o in od if o.get("status") in ("paid", "completed"))
        vis = len({v["visitor_id"] for v in visits if v["day"][5:7] == f"{m:02d}"})
        ws.append([MOIS[m - 1], dv, len(od), round(rev, 2), vis])
    ws.append(["Total", len(devis), len(orders),
               round(sum(float(o.get("total_amount") or 0) for o in orders if o.get("status") in ("paid", "completed")), 2),
               len({v["visitor_id"] for v in visits})])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    ws = wb.create_sheet("Devis")
    header(ws, ["Date", "Prénom", "Nom", "Courriel", "Téléphone", "Type de projet", "Services choisis",
                "Estimation calculateur ($)", "Statut", "Notes"], [12, 14, 16, 28, 14, 18, 50, 14, 12, 60])
    for d in devis:
        dt = _as_date(d.get("created_at"))
        calc = d.get("calculateur") or {}
        ws.append([dt.strftime("%Y-%m-%d") if dt else "", d.get("prenom", ""), d.get("nom", ""), d.get("email", ""),
                   d.get("telephone", ""), d.get("project_type", ""), ", ".join(d.get("plans_noms") or d.get("plans_choisis") or []),
                   calc.get("estimation") or "", d.get("status", ""), (d.get("notes") or "")[:2000]])

    ws = wb.create_sheet("Commandes")
    header(ws, ["Date", "Numéro", "Modèle", "Client", "Courriel", "Sous-total ($)", "Taxes ($)", "Total ($)", "Statut"],
           [12, 18, 28, 22, 28, 14, 12, 12, 12])
    for o in orders:
        dt = _as_date(o.get("created_at"))
        ws.append([dt.strftime("%Y-%m-%d") if dt else "", o.get("order_number", ""), o.get("kit_name", ""),
                   o.get("customer_name", ""), o.get("customer_email", ""), o.get("subtotal") or 0,
                   o.get("tax_amount") or 0, o.get("total_amount") or 0, o.get("status", "")])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Abrisia-rapport-{year}.xlsx"'},
    )
